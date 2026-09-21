"""
Build the TownHall BAR TASTE PLATE workbook (Spring/Summer 2026).

Rebuild after editing the lists below:
    python3 bar-taste-plate/build_taste_plate.py
    python3 <xlsx-skill>/scripts/recalc.py bar-taste-plate/BAR_TASTE_PLATE_SPRING-SUMMER-2026.xlsx

Everything a bartender/barback fills in is a yellow cell. NEED columns are
formulas (PAR minus ON HAND) - never type in them.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.comments import Comment

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "BAR_TASTE_PLATE_SPRING-SUMMER-2026.xlsx")

# ---------------------------------------------------------------- content ---
# PAR values carried over from the original BAR_TASTE_PLATE_SPRING-SUMMER2026.xlsx.
# (par_number, unit_label) -> unit shows in the cell, the number stays math-able.

BATCHES = [
    ("Green Goddess Biz",            8, "BTL"),
    ("Spicy Cucumber Biz",           8, "BTL"),
    ("Say Less",                     6, "BTL"),
    ("Light My Fire 2.0 Biz",        6, "BTL"),
    ("Turmeric Juice (LMF 2.0)",     6, "BTL"),
    ("No New Friends Biz",           6, "BTL"),
    ("Hot Girl Simmer Biz",          6, "BTL"),
    ("Alter Ego Biz",                6, "BTL"),
    ("Coconut Blue Spirulina",       6, "BTL"),
    ("Not Your Average Spritz",      6, "BTL"),
    ("Strawberry Allulose",          6, "BTL"),
    ("Tart Cherry Allulose",         3, "BTL"),
    ("Nitro Espresso Martini",       1, "CAMBRO"),
]
BATCH_NOTE = "Turmeric Juice is for Light My Fire 2.0 - this is NOT the wellness shot."

MIXERS = [
    ("Pineapple Juice",   4, "BTL"),
    ("Orange Juice",      4, "BTL"),
    ("Grapefruit Juice",  4, "BTL"),
    ("Watermelon Juice",  6, "BTL"),
    ("House Grenadine",   2, "BTL"),
    ("Olive Juice",       1, "BTL"),
    ("Simple Syrup",      6, "BTL"),
    ("Lemon Juice",       6, "BTL"),
    ("Lime Juice",        6, "BTL"),
    ("Bloody Mary Mix",   6, "BTL"),
    ("Espresso",          4, "BTL"),
]

FRESH_COUNT = [   # counted items (have a number par)
    ("Limes",   8, "PANS"),
    ("Lemons",  8, "PANS"),
]
FRESH_YN = [      # yes / low / 86 items
    "Olives", "Mint", "Parsley", "Basil", "Tarragon",
    "Dehydrated Strawberries", "Dried Calendula Flowers", "Thai Chili",
    "Hibiscus Leaves", "Cherries", "Espresso Beans", "Dried Dragonfruit",
    "Dragon Fruit Powder", "Electrolyte Salt", "Jalapenos", "Salt", "Sugar",
    "Angostura Bitters",
]

# Bar tools - PARs intentionally left blank; set them once with the bar lead.
TOOLS = [
    "Jiggers", "Stir Spoons", "Hand Strainers", "Muddlers",
    "Scraper", "Strainers", "Double Strainers",
]

SETUP = [
    ("Glasses",          "Spot check - clean, no chips"),
    ("Coolers",          "Stocked + clean"),
    ("Silverware",       "Stocked / full"),
    ("Napkins",          "Stocked"),
    ("Side Plates",      "Clean + stocked"),
    ("Sriracha",         "Wiped clean, enough on hand"),
    ("Salt and Pepper",  "Full + wiped clean"),
    ("Register Paper",   "Back-up at each printer"),
    ("Ice Scoops",       "In place"),
    ("Shaker Tins",      "In place at every well"),
    ("Pens",             "Stocked"),
    ("Crowlers",         "Labeled + stocked"),
    ("Crowler Machine",  "Cleaned"),
    ("Sharpies",         "By crowler machine"),
    ("Wine Pourer",      "In place"),
    ("Wine Opener",      "In place"),
]

# ----------------------------------------------------------------- styling ---
FONT = "Arial"
NAVY = "1F3A5F"
TEAL = "2E6F6A"
AMBER = "FFF2CC"
AMBER_STRONG = "FFD966"
RED_FILL = "F8CBAD"
GREY = "F2F2F2"
INPUT_FILL = PatternFill("solid", fgColor="FFFDE7")   # every fill-in cell
HDR_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor=TEAL)
BAND_FILL = PatternFill("solid", fgColor=GREY)

thin = Side(style="thin", color="BFBFBF")
med = Side(style="medium", color=NAVY)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def f(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)


def page(ws, landscape=False):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.45


def title_block(ws, subtitle, step, last_col="F", first_sheet=False, stacked=False):
    """Title + section name + the date/shift/name line. Returns the next free row."""
    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value = "TOWNHALL  |  BAR TASTE PLATE  -  SPRING / SUMMER 2026"
    c.font = f(15, True, "FFFFFF")
    c.fill = HDR_FILL
    c.alignment = CENTER
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f"A2:{last_col}2")
    c = ws["A2"]
    c.value = f"{step}   {subtitle}"
    c.font = f(11, True, "FFFFFF")
    c.fill = SUB_FILL
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 22

    # Date / shift / name. Entered once on page 1, pulled through on pages 2-4.
    fields = ["DATE", "SHIFT (AM/PM)", "COMPLETED BY"]
    srcs = ["$B$3", "$D$3", "$F$3"]
    if stacked:                      # narrow sheets: one field per row
        slots = [("A3", "B3"), ("A4", "B4"), ("A5", "B5")]
    else:
        slots = [("A3", "B3"), ("C3", "D3"), ("E3", "F3")]

    for i, (lbl_ref, val_ref) in enumerate(slots):
        lb = ws[lbl_ref]
        lb.value = fields[i]
        lb.font = f(9, True, NAVY)
        lb.alignment = Alignment(horizontal="right", vertical="center")
        cell = ws[val_ref]
        if not first_sheet:
            cell.value = (f"=IF('1. Batches & Mixers'!{srcs[i]}=\"\",\"\","
                          f"'1. Batches & Mixers'!{srcs[i]})")
        cell.font = f(10, True)
        cell.fill = INPUT_FILL
        cell.border = Border(bottom=Side(style="medium", color=NAVY))
        cell.alignment = CENTER
        if stacked:
            ws.merge_cells(start_row=cell.row, start_column=2, end_row=cell.row, end_column=3)
        ws.row_dimensions[cell.row].height = 22

    spacer = 6 if stacked else 4
    ws.row_dimensions[spacer].height = 6
    return spacer + 1


def legend(ws, row, text, last_col="F"):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws[f"A{row}"]
    c.value = text
    c.font = f(9, italic=True, color="595959")
    c.alignment = WRAP
    ws.row_dimensions[row].height = 26
    return row + 1


def table_header(ws, row, cols):
    for i, (label, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = f(9, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=thin, right=thin, top=med, bottom=med)
    ws.row_dimensions[row].height = 30
    return row + 1


def band(ws, row, text, last_col="F"):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws[f"A{row}"]
    c.value = text
    c.font = f(10, True, NAVY)
    c.fill = BAND_FILL
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c.border = Border(top=med, bottom=thin)
    last_idx = ord(last_col) - ord("A") + 1
    for col in range(2, last_idx + 1):
        ws.cell(row=row, column=col).fill = BAND_FILL
        ws.cell(row=row, column=col).border = Border(top=med, bottom=thin)
    ws.row_dimensions[row].height = 20
    return row + 1


def note(ws, row, text, last_col="F"):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws[f"A{row}"]
    c.value = text
    c.font = f(9, True, "9C0006")
    c.alignment = WRAP
    ws.row_dimensions[row].height = 18
    return row + 1


def count_row(ws, row, name, par, unit, example=False):
    """ITEM | PAR | ON HAND | NEED | TASTED & DATED | POSTED IN GM"""
    it = ws.cell(row=row, column=1, value=name)
    it.font = f(10, italic=example, color="808080" if example else "000000")
    it.alignment = LEFT

    p = ws.cell(row=row, column=2, value=par)
    p.number_format = f'0" {unit}"'
    p.font = f(10, bold=not example, color="808080" if example else "0000FF")
    p.alignment = CENTER

    oh = ws.cell(row=row, column=3)
    oh.font = f(10, bold=True, color="808080" if example else "000000")
    oh.alignment = CENTER
    if example:
        oh.value = 5
    else:
        oh.fill = INPUT_FILL

    nd = ws.cell(row=row, column=4, value=f'=IF(C{row}="","",MAX(0,B{row}-C{row}))')
    nd.font = f(10, True)
    nd.alignment = CENTER

    for col in (5, 6):
        c = ws.cell(row=row, column=col)
        c.alignment = CENTER
        c.font = f(10, color="808080" if example else "000000")
        if not example:
            c.fill = INPUT_FILL
    if example:
        ws.cell(row=row, column=5).value = "OK"
        ws.cell(row=row, column=6).value = "N/A"

    for col in range(1, 7):
        cell = ws.cell(row=row, column=col)
        cell.border = BOX
        if example:
            cell.fill = PatternFill("solid", fgColor="F2F2F2")
    ws.row_dimensions[row].height = 20
    return row + 1


def add_dv(ws, options, cells, prompt=None):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(options),
                        allow_blank=True, showDropDown=False)
    dv.error = "Pick one from the drop-down."
    dv.errorTitle = "Not on the list"
    if prompt:
        dv.prompt = prompt
        dv.promptTitle = "Fill this in"
    ws.add_data_validation(dv)
    for rng in cells:
        dv.add(rng)


def need_formatting(ws, rng_need, rng_status):
    """NEED > 0 lights up amber; an 86 lights up red."""
    ws.conditional_formatting.add(
        rng_need,
        CellIsRule(operator="greaterThan", formula=["0"],
                   fill=PatternFill("solid", fgColor=AMBER_STRONG),
                   font=Font(name=FONT, size=10, bold=True, color="9C0006")))
    if rng_status:
        ws.conditional_formatting.add(
            rng_status,
            CellIsRule(operator="equal", formula=['"86"'],
                       fill=PatternFill("solid", fgColor=RED_FILL),
                       font=Font(name=FONT, size=10, bold=True, color="9C0006")))


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


COUNT_COLS = [("ITEM", 34), ("PAR", 11), ("ON HAND", 11), ("NEED", 10),
              ("TASTED &\nDATED", 13), ("POSTED IN GM /\nPREP TOLD", 16)]

wb = Workbook()

# =========================================================== 1. BATCHES =====
ws = wb.active
ws.title = "1. Batches & Mixers"
page(ws)
widths(ws, {"A": 34, "B": 11, "C": 11, "D": 10, "E": 13, "F": 16})
r = title_block(ws, "BATCHES, COCKTAIL PRODUCTS & MIXERS", "STEP 1 of 4", first_sheet=True)
r = legend(ws, r,
           "Fill in the YELLOW cells only. NEED does the math for you (PAR minus ON HAND) and turns "
           "gold when you have to make or order more. TASTED & DATED: OK / LOW / 86.")
r = table_header(ws, r, COUNT_COLS)
example_row = r
r = count_row(ws, r, "EXAMPLE - do not count this row", 8, "BTL", example=True)

r = band(ws, r, "COCKTAIL PRODUCTS  /  BATCHES")
batch_start = r
for name, par, unit in BATCHES:
    r = count_row(ws, r, name, par, unit)
batch_end = r - 1
r = note(ws, r, "* " + BATCH_NOTE)

r = band(ws, r, "JUICES & MIXERS")
mix_start = r
for name, par, unit in MIXERS:
    r = count_row(ws, r, name, par, unit)
mix_end = r - 1

r += 1
r = note(ws, r, "86'd ANYTHING? Tell a manager AND post it in GroupMe before you leave the bar.")
ws.freeze_panes = "A%d" % (example_row + 1)
ws.print_area = f"A1:F{r}"

add_dv(ws, ["OK", "LOW", "86"], [f"E{batch_start}:E{batch_end}", f"E{mix_start}:E{mix_end}"])
add_dv(ws, ["YES", "NO", "N/A"], [f"F{batch_start}:F{batch_end}", f"F{mix_start}:F{mix_end}"])
need_formatting(ws, f"D{batch_start}:D{mix_end}", f"E{batch_start}:E{mix_end}")
ws["B%d" % batch_start].comment = Comment(
    "PAR values carried over from the original Spring/Summer 2026 taste plate. "
    "Change a PAR here and the NEED column updates itself.", "TownHall")

# ============================================================= 2. FRESH =====
ws = wb.create_sheet("2. Fresh & Garnish")
page(ws)
widths(ws, {"A": 34, "B": 11, "C": 11, "D": 10, "E": 13, "F": 16})
r = title_block(ws, "FRESH FRUIT, HERBS & GARNISH", "STEP 2 of 4")
r = legend(ws, r,
           "BARBACKS FILL THIS PAGE OUT. Limes and lemons get counted in sixth pans. Everything "
           "else is just a status: OK / LOW / 86. Anything LOW or 86 - tell a manager and post in GroupMe.")
r = table_header(ws, r, [("ITEM", 34), ("PAR", 11), ("ON HAND", 11), ("NEED", 10),
                         ("STATUS", 13), ("POSTED IN GM /\nPREP TOLD", 16)])
r = band(ws, r, "COUNTED - SIXTH PANS")
fc_start = r
for name, par, unit in FRESH_COUNT:
    r = count_row(ws, r, name, par, unit)
fc_end = r - 1

r = band(ws, r, "GARNISH & DRY GOODS  -  OK / LOW / 86")
yn_start = r
for name in FRESH_YN:
    it = ws.cell(row=r, column=1, value=name)
    it.font = f(10)
    it.alignment = LEFT
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    mid = ws.cell(row=r, column=2, value="-")
    mid.alignment = CENTER
    mid.font = f(9, color="A6A6A6")
    for col in (5, 6):
        c = ws.cell(row=r, column=col)
        c.fill = INPUT_FILL
        c.alignment = CENTER
        c.font = f(10)
    for col in range(1, 7):
        ws.cell(row=r, column=col).border = BOX
    ws.row_dimensions[r].height = 20
    r += 1
yn_end = r - 1

r += 1
r = note(ws, r, "86'd ANYTHING? Tell a manager AND post it in GroupMe before you leave the bar.")
ws.freeze_panes = "A7"
ws.print_area = f"A1:F{r}"
add_dv(ws, ["OK", "LOW", "86"], [f"E{fc_start}:E{yn_end}"])
add_dv(ws, ["YES", "NO", "N/A"], [f"F{fc_start}:F{yn_end}"])
need_formatting(ws, f"D{fc_start}:D{fc_end}", f"E{fc_start}:E{yn_end}")

# ============================================================= 3. TOOLS =====
ws = wb.create_sheet("3. Bar Tools")
page(ws)
widths(ws, {"A": 26, "B": 11, "C": 11, "D": 10, "E": 13, "F": 24})
r = title_block(ws, "BAR TOOL COUNT", "STEP 3 of 4")
r = legend(ws, r,
           "Count every well plus the back-up drawer. PAR is blank on purpose - set it once with "
           "your bar lead and it stays put. NEED fills itself in. Use NOTES for anything broken, "
           "bent, or walked off.")
r = table_header(ws, r, [("TOOL", 26), ("PAR", 11), ("ON HAND", 11), ("NEED", 10),
                         ("CONDITION", 13), ("NOTES  (broken / missing / where)", 24)])
tool_start = r
for name in TOOLS:
    it = ws.cell(row=r, column=1, value=name)
    it.font = f(10)
    it.alignment = LEFT
    p = ws.cell(row=r, column=2)
    p.fill = INPUT_FILL
    p.font = f(10, True, "0000FF")
    p.alignment = CENTER
    oh = ws.cell(row=r, column=3)
    oh.fill = INPUT_FILL
    oh.font = f(10, True)
    oh.alignment = CENTER
    nd = ws.cell(row=r, column=4, value=f'=IF(OR(B{r}="",C{r}=""),"",MAX(0,B{r}-C{r}))')
    nd.font = f(10, True)
    nd.alignment = CENTER
    for col in (5, 6):
        c = ws.cell(row=r, column=col)
        c.fill = INPUT_FILL
        c.alignment = CENTER if col == 5 else LEFT
        c.font = f(10)
    for col in range(1, 7):
        ws.cell(row=r, column=col).border = BOX
    ws.row_dimensions[r].height = 24
    r += 1
tool_end = r - 1

r += 1
tot_lbl = ws.cell(row=r, column=1, value="TOOLS SHORT TODAY")
tot_lbl.font = f(10, True, NAVY)
tot_lbl.alignment = LEFT
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
tot = ws.cell(row=r, column=4, value=f"=SUM(D{tool_start}:D{tool_end})")
tot.font = f(11, True, "9C0006")
tot.alignment = CENTER
tot.border = BOX
ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
ws.cell(row=r, column=5, value="Anything short or broken -> tell a manager").font = f(9, italic=True, color="595959")
ws.cell(row=r, column=5).alignment = LEFT
r += 2
r = note(ws, r, "Tools go in the dish pit dirty, not in the trash. A missing jigger is a comp'd drink.")
ws.freeze_panes = "A7"
ws.print_area = f"A1:F{r}"
add_dv(ws, ["OK", "DIRTY", "BROKEN", "MISSING"], [f"E{tool_start}:E{tool_end}"])
need_formatting(ws, f"D{tool_start}:D{tool_end}", None)
ws.conditional_formatting.add(
    f"E{tool_start}:E{tool_end}",
    FormulaRule(formula=[f'OR($E{tool_start}="BROKEN",$E{tool_start}="MISSING")'],
                fill=PatternFill("solid", fgColor=RED_FILL),
                font=Font(name=FONT, size=10, bold=True, color="9C0006")))
ws["B%d" % tool_start].comment = Comment(
    "Bar tool PARs were not on the original sheet. Set them once (per well + back-ups) "
    "and the NEED column takes care of the rest.", "TownHall")

# ==================================================== 4. SETUP & COMMS ======
ws = wb.create_sheet("4. Setup & Comms")
page(ws)
widths(ws, {"A": 24, "B": 34, "C": 10})
r = title_block(ws, "BAR SET-UP CHECK + DAILY COMMUNICATION", "STEP 4 of 4",
                last_col="C", stacked=True)
r = legend(ws, r, "Walk the bar, check each line off, then finish the communication box at the bottom.",
           last_col="C")
r = table_header(ws, r, [("ITEM", 24), ("WHAT TO CHECK", 34), ("DONE", 10)])
setup_start = r
for name, action in SETUP:
    ws.cell(row=r, column=1, value=name).font = f(10)
    ws.cell(row=r, column=1).alignment = LEFT
    ws.cell(row=r, column=2, value=action).font = f(10, color="404040")
    ws.cell(row=r, column=2).alignment = LEFT
    c = ws.cell(row=r, column=3)
    c.fill = INPUT_FILL
    c.alignment = CENTER
    c.font = f(10, True)
    for col in range(1, 4):
        ws.cell(row=r, column=col).border = BOX
    ws.row_dimensions[r].height = 19
    r += 1
setup_end = r - 1
add_dv(ws, ["YES", "NO"], [f"C{setup_start}:C{setup_end}"])
ws.conditional_formatting.add(
    f"C{setup_start}:C{setup_end}",
    CellIsRule(operator="equal", formula=['"NO"'],
               fill=PatternFill("solid", fgColor=RED_FILL),
               font=Font(name=FONT, size=10, bold=True, color="9C0006")))

r += 1
r = band(ws, r, "DAILY COMMUNICATION", last_col="C")
COMMS = [
    ("New drafts poured through at every beer tower?", ["YES", "NO"]),
    ("8 oz pour beer taps marked?", ["YES", "NO"]),
    ("Anything 86'd? (list it below + post in GroupMe)", ["NO", "YES"]),
]
comm_rows = []
for q, opts in COMMS:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = ws.cell(row=r, column=1, value=q)
    c.font = f(10, True)
    c.alignment = LEFT
    c.border = BOX
    ws.cell(row=r, column=2).border = BOX
    a = ws.cell(row=r, column=3)
    a.fill = INPUT_FILL
    a.alignment = CENTER
    a.font = f(10, True)
    a.border = BOX
    add_dv(ws, opts, [f"C{r}"])
    comm_rows.append(r)
    ws.row_dimensions[r].height = 20
    r += 1
ws.conditional_formatting.add(
    f"C{comm_rows[0]}:C{comm_rows[1]}",
    CellIsRule(operator="equal", formula=['"NO"'],
               fill=PatternFill("solid", fgColor=RED_FILL),
               font=Font(name=FONT, size=10, bold=True, color="9C0006")))
ws.conditional_formatting.add(
    f"C{comm_rows[2]}",
    CellIsRule(operator="equal", formula=['"YES"'],
               fill=PatternFill("solid", fgColor=AMBER_STRONG),
               font=Font(name=FONT, size=10, bold=True, color="9C0006")))

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
c = ws.cell(row=r, column=1, value="86'd ITEMS / NOTES FOR THE NEXT SHIFT")
c.font = f(10, True, NAVY)
c.alignment = LEFT
c.fill = BAND_FILL
r += 1
for _ in range(3):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(row=r, column=1)
    c.fill = INPUT_FILL
    c.border = Border(bottom=thin, left=thin, right=thin, top=thin)
    ws.row_dimensions[r].height = 22
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
c = ws.cell(row=r, column=1, value="SIGN OFF")
c.font = f(10, True, "FFFFFF")
c.fill = SUB_FILL
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
r += 1
for label in ("Bartender / barback", "Time finished", "Manager verified"):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    lb = ws.cell(row=r, column=1, value=label)
    lb.font = f(10, True)
    lb.alignment = LEFT
    lb.border = BOX
    ws.cell(row=r, column=2).border = BOX
    sg = ws.cell(row=r, column=3)
    sg.fill = INPUT_FILL
    sg.border = BOX
    ws.row_dimensions[r].height = 22
    r += 1

r += 1
r = note(ws, r, "A manager must be told about every 86'd item AND it gets posted in GroupMe.", last_col="C")
ws.print_area = f"A1:C{r}"

wb.save(OUT)
print("wrote", OUT)
